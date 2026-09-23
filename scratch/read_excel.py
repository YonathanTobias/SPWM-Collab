import os
import zipfile
import xml.etree.ElementTree as ET

excel_path = r"C:\Users\Tobias\Downloads\MATRIX KERJASAMA TRIWULAN II 2026.xlsx"

print("Excel File Exists:", os.path.exists(excel_path))

try:
    import pandas as pd
    df = pd.read_excel(excel_path)
    print("Read with pandas successfully! Columns:", df.columns.tolist())
    print("First 5 rows:")
    print(df.head())
    df.to_csv("excel_extracted.csv", index=False)
except Exception as e:
    print("Pandas failed:", e)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path)
        print("Sheets:", wb.sheetnames)
        sheet = wb.active
        for row in list(sheet.iter_rows(values_only=True))[:10]:
            print(row)
    except Exception as e2:
        print("Openpyxl failed:", e2)
        # Zipfile fallback for xlsx
        with zipfile.ZipFile(excel_path, 'r') as z:
            print("Zip files inside xlsx:", [f for f in z.namelist() if f.startswith('xl/')])
